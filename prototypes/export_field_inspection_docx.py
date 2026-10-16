"""
Engine xuất Biên bản Kiểm tra, Khảo sát và Nghiệm thu Hiện trường (.docx)
tuân thủ Nghị định số 30/2020/NĐ-CP và Nghị định số 207/2026/NĐ-CP:
- 100% phông chữ Times New Roman, màu chữ đen tuyền (#000000).
- Bảng biểu nền trắng tự nhiên, viền đơn đen mảnh 0.5pt (sz="4"), lặp lại tiêu đề bảng (tblHeader) và chống cắt dòng (cantSplit).
- Tuyệt đối không dùng ký tự xuống dòng mềm (\n) trong các đoạn văn căn đều hai bên (JUSTIFY).
"""

import sys
from pathlib import Path
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


BLACK = RGBColor(0, 0, 0)
FONT_NAME = "Times New Roman"


def set_run_font(run, size_pt: float, bold: bool = False, italic: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), FONT_NAME)
    rFonts.set(qn("w:hAnsi"), FONT_NAME)
    rFonts.set(qn("w:cs"), FONT_NAME)
    rFonts.set(qn("w:eastAsia"), FONT_NAME)


def add_formatted_paragraph(
    doc,
    text: str,
    size_pt: float = 13.0,
    bold: bool = False,
    italic: bool = False,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    space_before_pt: float = 3.0,
    space_after_pt: float = 3.0,
):
    # Đảm bảo không có ký tự xuống dòng mềm bên trong đoạn văn
    clean_text = text.replace("\n", " ").strip()
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(clean_text)
    set_run_font(run, size_pt=size_pt, bold=bold, italic=italic)
    return p


def set_table_borders_05pt_black(table) -> None:
    tblPr = table._tbl.tblPr
    tblBorders = tblPr.find(qn("w:tblBorders"))
    if tblBorders is not None:
        tblPr.remove(tblBorders)
    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")  # 0.5 pt
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def set_table_borderless(table) -> None:
    tblPr = table._tbl.tblPr
    tblBorders = tblPr.find(qn("w:tblBorders"))
    if tblBorders is not None:
        tblPr.remove(tblBorders)
    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "none")
        border.set(qn("w:sz"), "0")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def configure_row_Repeat_and_cant_split(row, is_header: bool = False) -> None:
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)
    if is_header:
        tblHeader = OxmlElement("w:tblHeader")
        trPr.append(tblHeader)


def generate_field_inspection_report(output_docx_path: Path) -> Path:
    doc = docx.Document()

    # Thiết lập lề trang chuẩn Nghị định số 30/2020/NĐ-CP: Trên 2cm, Dưới 2cm, Trái 3cm, Phải 1.5cm
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)

    # 1. Bảng Quốc hiệu & Tiêu ngữ (2 cột, không viền)
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borderless(header_table)

    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)

    p_org1 = cell_left.paragraphs[0]
    p_org1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p_org1.add_run("BAN QUẢN LÝ DỰ ÁN ĐẦU TƯ XÂY DỰNG"), 12.0, bold=False)

    p_org2 = cell_left.add_paragraph()
    p_org2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p_org2.add_run("ĐOÀN KIỂM TRA - GIÁM SÁT HIỆN TRƯỜNG"), 12.0, bold=True)

    p_org3 = cell_left.add_paragraph()
    p_org3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p_org3.add_run("Số: 01/2026/BBKT-HT"), 12.0, bold=False)

    p_qh1 = cell_right.paragraphs[0]
    p_qh1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p_qh1.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"), 12.0, bold=True)

    p_qh2 = cell_right.add_paragraph()
    p_qh2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p_qh2.add_run("Độc lập - Tự do - Hạnh phúc"), 13.0, bold=True)

    p_date = cell_right.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p_date.add_run("Hà Nội, ngày 16 tháng 10 năm 2026"), 12.0, italic=True)

    # 2. Tên Biên bản
    add_formatted_paragraph(
        doc,
        "BIÊN BẢN KIỂM TRA, KHẢO SÁT VÀ NGHIỆM THU CÔNG VIỆC XÂY DỰNG TẠI HIỆN TRƯỜNG",
        size_pt=14.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before_pt=12.0,
        space_after_pt=6.0,
    )

    # 3. Căn cứ pháp lý (mỗi dòng là một đoạn độc lập, căn trái)
    add_formatted_paragraph(
        doc,
        "I. CĂN CỨ PHÁP LÝ VÀ TIÊU CHUẨN KỸ THUẬT ÁP DỤNG",
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before_pt=6.0,
        space_after_pt=3.0,
    )

    legal_bases = [
        "- Căn cứ Luật Xây dựng số 135/2025/QH15;",
        "- Căn cứ Luật Phòng cháy, chữa cháy và cứu nạn, cứu hộ số 55/2024/QH15;",
        "- Căn cứ Nghị định số 207/2026/NĐ-CP quy định chi tiết về quản lý chất lượng, thi công xây dựng và bảo trì công trình xây dựng (Điều 22 về nghiệm thu công việc xây dựng);",
        "- Căn cứ Nghị định số 217/2026/NĐ-CP quy định chi tiết một số nội dung về quản lý dự án đầu tư xây dựng;",
        "- Căn cứ Nghị định số 347/2026/NĐ-CP quy định về công tác kiểm tra hiện trường và nghiệm thu về phòng cháy, chữa cháy;",
        "- Căn cứ Hồ sơ thiết kế bản vẽ thi công và Chỉ dẫn kỹ thuật đã được phê duyệt.",
    ]
    for item in legal_bases:
        add_formatted_paragraph(
            doc,
            item,
            size_pt=13.0,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            space_before_pt=2.0,
            space_after_pt=2.0,
        )

    # 4. Thời gian, địa điểm và thành phần tham gia
    add_formatted_paragraph(
        doc,
        "II. THỜI GIAN, ĐỊA ĐIỂM VÀ THÀNH PHẦN TRỰC TIẾP KIỂM TRA TẠI HIỆN TRƯỜNG",
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before_pt=8.0,
        space_after_pt=3.0,
    )
    add_formatted_paragraph(
        doc,
        "- Thời gian bắt đầu: 08 giờ 30 phút, ngày 16 tháng 10 năm 2026; kết thúc lúc 11 giờ 30 phút cùng ngày.",
        size_pt=13.0,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_formatted_paragraph(
        doc,
        "- Địa điểm kiểm tra tại hiện trường: Khu vực Tầng hầm B1, Trục (A-D)/(1-6), Tòa nhà Trung tâm Nghiên cứu và Phát triển Công nghệ.",
        size_pt=13.0,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_formatted_paragraph(
        doc,
        "1. Đại diện Chủ đầu tư và Bộ phận Giám sát thi công xây dựng:",
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_formatted_paragraph(
        doc,
        "- Ông/Bà: Nguyễn Văn A - Chức vụ: Giám sát trưởng thi công xây dựng.",
        size_pt=13.0,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_formatted_paragraph(
        doc,
        "2. Đại diện Nhà thầu khảo sát và thi công xây dựng:",
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_formatted_paragraph(
        doc,
        "- Ông/Bà: Trần Văn B - Chức vụ: Chỉ huy trưởng công trường / Phụ trách kỹ thuật thi công trực tiếp.",
        size_pt=13.0,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    # 5. Bảng kết quả kiểm tra, đo đạc thực tế tại hiện trường
    add_formatted_paragraph(
        doc,
        "III. KẾT QUẢ KIỂM TRA, ĐO ĐẠC THỰC TẾ TẠI HIỆN TRƯỜNG",
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before_pt=8.0,
        space_after_pt=4.0,
    )

    headers = [
        "TT",
        "Nội dung kiểm tra tại hiện trường",
        "Tiêu chuẩn / Thiết kế yêu cầu",
        "Số liệu thực đo tại hiện trường",
        "Đánh giá",
    ]
    rows_data = [
        (
            "1",
            "Kiểm tra hệ thống lưới trắc địa và cao độ sàn tầng hầm B1 tại trục (A-D)/(1-6)",
            "TCVN 9398:2012; sai lệch tim trục không quá ±5 mm",
            "Sai lệch thực đo bằng máy toàn đạc điện tử: +2.0 mm",
            "Đạt",
        ),
        (
            "2",
            "Kiểm tra lắp dựng cốt thép dầm sàn tầng hầm B1 và chiều dày lớp bê tông bảo vệ",
            "TCVN 4453:1995; thép D25 CB500-V; lớp bảo vệ dầm a = 25 mm (±5 mm)",
            "Đúng chủng loại thép D25; chiều dày lớp bảo vệ thực đo từ 25 mm đến 27 mm",
            "Đạt",
        ),
        (
            "3",
            "Kiểm tra độ sụt hỗn hợp bê tông thương phẩm B30 và công tác lấy mẫu tại hiện trường",
            "TCVN 3106:2022; độ sụt 160 ± 20 mm; lấy 01 tổ mẫu / 50 m3",
            "Độ sụt thực đo: 165 mm; đã lấy 03 tổ mẫu lập phương có ký nhận tại chỗ",
            "Đạt",
        ),
        (
            "4",
            "Kiểm tra ống chờ xuyên vách và chi tiết chèn bịt ngăn cháy trục kỹ thuật cơ điện",
            "QCVN 06:2022/BXD; Nghị định số 347/2026/NĐ-CP; giới hạn chịu lửa EI 60",
            "Ống lồng thép đặt đúng tọa độ thiết kế, vật liệu chèn ngăn cháy có hồ sơ kiểm định",
            "Đạt",
        ),
    ]

    table = doc.add_table(rows=1 + len(rows_data), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_05pt_black(table)

    col_widths_cm = [1.2, 5.2, 4.2, 4.3, 1.6]

    # Hàng tiêu đề bảng
    hdr_row = table.rows[0]
    configure_row_Repeat_and_cant_split(hdr_row, is_header=True)
    for col_idx, text in enumerate(headers):
        cell = hdr_row.cells[col_idx]
        cell.width = Cm(col_widths_cm[col_idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3.0)
        p.paragraph_format.space_after = Pt(3.0)
        set_run_font(p.add_run(text), size_pt=11.5, bold=True)

    # Các hàng dữ liệu
    for row_idx, rdata in enumerate(rows_data, start=1):
        row = table.rows[row_idx]
        configure_row_Repeat_and_cant_split(row, is_header=False)
        for col_idx, val in enumerate(rdata):
            cell = row.cells[col_idx]
            cell.width = Cm(col_widths_cm[col_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            if col_idx in (0, 4):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                is_bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                is_bold = False
            p.paragraph_format.space_before = Pt(3.0)
            p.paragraph_format.space_after = Pt(3.0)
            set_run_font(p.add_run(val), size_pt=11.5, bold=is_bold)

    # 6. Kết luận nghiệm thu theo Điều 22 Nghị định số 207/2026/NĐ-CP
    add_formatted_paragraph(
        doc,
        "IV. KẾT LUẬN CỦA ĐOÀN KIỂM TRA VÀ GIÁM SÁT HIỆN TRƯỜNG",
        size_pt=13.0,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before_pt=8.0,
        space_after_pt=3.0,
    )
    add_formatted_paragraph(
        doc,
        "- Các công việc khảo sát trắc địa, lắp dựng cốt thép và đặt ống chờ kỹ thuật tại khu vực Tầng hầm B1 trục (A-D)/(1-6) đã được thực hiện phù hợp với hồ sơ thiết kế bản vẽ thi công được duyệt và các quy chuẩn, tiêu chuẩn kỹ thuật hiện hành.",
        size_pt=13.0,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_formatted_paragraph(
        doc,
        "- Đồng ý nghiệm thu công việc xây dựng theo quy định tại Điều 22 Nghị định số 207/2026/NĐ-CP và cho phép Nhà thầu thi công chuyển bước thực hiện công tác tiếp theo.",
        size_pt=13.0,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )

    # 7. Bảng chữ ký xác nhận (2 cột, không viền)
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borderless(sig_table)

    sig_left = sig_table.cell(0, 0)
    sig_right = sig_table.cell(0, 1)

    p_sl1 = sig_left.paragraphs[0]
    p_sl1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sl1.paragraph_format.space_before = Pt(10.0)
    set_run_font(p_sl1.add_run("NGƯỜI PHỤ TRÁCH KỸ THUẬT THI CÔNG TRỰC TIẾP"), 12.0, bold=True)

    p_sl2 = sig_left.add_paragraph()
    p_sl2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p_sl2.add_run("(Ký, ghi rõ họ tên)"), 11.5, italic=True)

    p_sr1 = sig_right.paragraphs[0]
    p_sr1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sr1.paragraph_format.space_before = Pt(10.0)
    set_run_font(p_sr1.add_run("NGƯỜI GIÁM SÁT THI CÔNG XÂY DỰNG"), 12.0, bold=True)

    p_sr2 = sig_right.add_paragraph()
    p_sr2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p_sr2.add_run("(Ký, ghi rõ họ tên)"), 11.5, italic=True)

    output_docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_docx_path)
    return output_docx_path


if __name__ == "__main__":
    target_docx = (
        Path(__file__).resolve().parent.parent
        / "output_reports"
        / "Bien_Ban_Kiem_Tra_Khao_Sat_Va_Nghiem_Thu_Hien_Truong.docx"
    )
    saved = generate_field_inspection_report(target_docx)
    print(f"Đã tạo Biên bản Kiểm tra, Khảo sát và Nghiệm thu Hiện trường (.docx) tại: {saved}")
