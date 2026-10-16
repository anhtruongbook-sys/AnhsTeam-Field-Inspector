"""
Công cụ xuất Bảng kiểm tra hiện trường (Field Inspection & Survey Checklist .xlsx)
sử dụng 100% công thức Excel động Native (COUNTIF, COUNTA, IF, ROUND) liên kết giữa các ô.
Tuân thủ Luật Xây dựng số 135/2025/QH15, Nghị định số 207/2026/NĐ-CP và Nghị định số 217/2026/NĐ-CP.
"""

import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def create_dynamic_field_checklist(output_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Checklist_Hien_Truong"

    font_family = "Times New Roman"
    title_font = Font(name=font_family, size=13, bold=True, color="000000")
    subtitle_font = Font(name=font_family, size=11, italic=True, color="000000")
    header_font = Font(name=font_family, size=11, bold=True, color="000000")
    bold_font = Font(name=font_family, size=11, bold=True, color="000000")
    regular_font = Font(name=font_family, size=11, bold=False, color="000000")

    thin_side = Side(border_style="thin", color="000000")
    grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    white_fill = PatternFill(fill_type=None)

    # Tiêu đề bảng kiểm tra
    ws.merge_cells("A1:I1")
    ws["A1"] = "BẢNG KIỂM TRA, KHẢO SÁT VÀ GIÁM SÁT CHẤT LƯỢNG TẠI HIỆN TRƯỜNG"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:I2")
    ws["A2"] = (
        "Căn cứ pháp lý: Luật Xây dựng số 135/2025/QH15; Nghị định số 207/2026/NĐ-CP; "
        "Nghị định số 217/2026/NĐ-CP; Nghị định số 347/2026/NĐ-CP"
    )
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Khu vực thông số tổng hợp tự động bằng công thức Excel động (Dòng 4 - 6)
    ws["B4"] = "Tổng số chỉ tiêu kiểm tra:"
    ws["B4"].font = bold_font
    ws["C4"] = "=COUNTA(B10:B19)"
    ws["C4"].font = bold_font
    ws["C4"].alignment = Alignment(horizontal="center")

    ws["D4"] = "Số chỉ tiêu ĐẠT (1):"
    ws["D4"].font = bold_font
    ws["E4"] = '=COUNTIF(G10:G19, 1)'
    ws["E4"].font = bold_font
    ws["E4"].alignment = Alignment(horizontal="center")

    ws["F4"] = "Số chỉ tiêu KHÔNG ĐẠT (0):"
    ws["F4"].font = bold_font
    ws["G4"] = '=COUNTIF(G10:G19, 0)'
    ws["G4"].font = bold_font
    ws["G4"].alignment = Alignment(horizontal="center")

    ws["B5"] = "Tỷ lệ tuân thủ kỹ thuật (%):"
    ws["B5"].font = bold_font
    ws["C5"] = "=IF(C4>0, ROUND(E4/C4*100, 1), 0)"
    ws["C5"].font = bold_font
    ws["C5"].alignment = Alignment(horizontal="center")

    ws["D5"] = "Kết luận đủ điều kiện nghiệm thu:"
    ws["D5"].font = bold_font
    ws.merge_cells("E5:G5")
    ws["E5"] = '=IF(G4=0, "ĐỦ ĐIỀU KIỆN CHUYỂN BƯỚC / NGHIỆM THU", "YÊU CẦU KHẮC PHỤC SAI SÓT (NCR)")'
    ws["E5"].font = bold_font
    ws["E5"].alignment = Alignment(horizontal="center")

    for r in (4, 5):
        for c in range(2, 8):
            cell = ws.cell(row=r, column=c)
            cell.border = grid_border
            cell.fill = white_fill

    # Dòng tiêu đề bảng dữ liệu (Dòng 9)
    headers = [
        "TT",
        "Phân nhóm chuyên ngành",
        "Nội dung kiểm tra / khảo sát tại hiện trường",
        "Tiêu chuẩn / Quy chuẩn áp dụng",
        "Giá trị thiết kế / Yêu cầu",
        "Giá trị thực đo tại hiện trường",
        "Mã điểm đánh giá (1=Đạt, 0=K.Đạt)",
        "Kết luận tự động",
        "Ghi chú / Vị trí trục",
    ]

    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=9, column=col_idx, value=text)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = grid_border
        cell.fill = white_fill
    ws.row_dimensions[9].height = 32

    # Dữ liệu mẫu kiểm tra hiện trường (10 chỉ tiêu thực tế)
    sample_items = [
        (
            "Khảo sát & Trắc địa",
            "Kiểm tra hệ thống lưới khống chế trắc địa và mốc cao độ chuẩn tại hiện trường",
            "TCVN 9398:2012; NĐ 217/2026/NĐ-CP",
            "Sai số tim trục <= 5 mm; mốc cao độ ổn định",
            "Sai số đo kiểm chứng: +2.0 mm",
            1,
            "Trục (A-D)/(1-6)",
        ),
        (
            "Khảo sát Hiện trạng",
            "Quan trắc lún và kiểm tra vết nứt công trình lân cận trước khi đào hố móng",
            "TCVN 9360:2012; NĐ 207/2026/NĐ-CP",
            "Ghi nhận đầy đủ biên bản + ảnh GPS hiện trạng",
            "Đã lập hồ sơ 04 hộ liền kề, độ nghiêng 0.08%",
            1,
            "Ranh giới phía Đông & Tây",
        ),
        (
            "Kết cấu BTCT",
            "Kiểm tra đường kính, khoảng cách và chiều dài nối chồng cốt thép dầm sàn",
            "TCVN 4453:1995; TCVN 5574:2018",
            "Thép chủ D25 CB500-V, nối chồng L >= 40d (1000 mm)",
            "Đúng chủng loại D25 CB500-V, chiều dài nối 1020 mm",
            1,
            "Dầm B1-G12 Trục (B-C)/3",
        ),
        (
            "Kết cấu BTCT",
            "Kiểm tra chiều dày lớp bê tông bảo vệ cốt thép chịu lực bằng con kê định hình",
            "TCVN 4453:1995 (Bảng 6)",
            "Dầm a = 25 mm (±5 mm); Sàn a = 15 mm (±3 mm)",
            "Dầm đạt 25-27 mm; Sàn đạt 15-16 mm",
            1,
            "Sàn tầng hầm B1",
        ),
        (
            "Vật liệu hiện trường",
            "Kiểm tra độ sụt hỗn hợp bê tông thương phẩm B30 (M400) trước khi bơm",
            "TCVN 3106:2022; Điều 22 NĐ 207/2026/NĐ-CP",
            "Độ sụt thiết kế: 160 ± 20 mm; lấy 01 tổ mẫu / 50 m3",
            "Độ sụt thực đo: 165 mm; đã đúc 03 tổ mẫu",
            1,
            "Trạm trộn xe số 01-03",
        ),
        (
            "Kiểm định NDT",
            "Kiểm tra cường độ bê tông cột ở tuổi 28 ngày bằng súng bật nẩy & siêu âm",
            "TCVN 9334:2012; TCVN 9357:2012",
            "Cường độ hiện trường đạt cấp độ bền B30 (>= 38.5 MPa)",
            "Cường độ quy đổi trung bình: 41.2 MPa",
            1,
            "Cột C1-C6 Tầng hầm B1",
        ),
        (
            "Cơ điện (MEP)",
            "Thử áp lực thủy tĩnh đường ống cấp nước sinh hoạt và ống chữa cháy vách tường",
            "TCVN 4513:1988; TCVN 7336:2021",
            "Áp lực thử 1.0 MPa trong 60 phút, độ sụt áp = 0 MPa",
            "Duy trì 1.0 MPa trong 60 phút, không rò rỉ",
            1,
            "Trục đứng kỹ thuật Nước",
        ),
        (
            "Cơ điện (MEP)",
            "Đo điện trở nối đất hệ thống chống sét trực tiếp và tiếp địa an toàn điện",
            "TCVN 9385:2012; TCVN 7447",
            "Tiếp địa chống sét R <= 10 Ohm; An toàn điện R <= 4 Ohm",
            "Chống sét R = 2.4 Ohm; An toàn điện R = 1.1 Ohm",
            1,
            "Hố kiểm tra tiếp địa số 1-4",
        ),
        (
            "PCCC & Thoát nạn",
            "Kiểm tra kích thước thông thủy lối thoát nạn và cửa chống cháy buồng thang bộ",
            "Luật 55/2024/QH15; QCVN 06:2022/BXD",
            "Rộng thông thủy cửa >= 0.90 m, giới hạn chịu lửa EI 60",
            "Rộng thông thủy thực đo 0.95 m, đầy đủ tem KĐ PCCC",
            1,
            "Buồng thang thoát nạn ST1, ST2",
        ),
        (
            "PCCC Liên động",
            "Thử nghiệm chênh áp buồng thang bộ N1/N3 và hút khói hành lang khi có tín hiệu báo cháy",
            "QCVN 06:2022/BXD; NĐ 347/2026/NĐ-CP",
            "Chênh áp buồng thang 20 - 50 Pa; liên động ngắt nguồn",
            "Chênh áp thực đo đạt 34 Pa; quạt khởi động sau 5 giây",
            1,
            "Toàn bộ trục thang bộ thoát nạn",
        ),
    ]

    for idx, item in enumerate(sample_items, start=1):
        row_num = 9 + idx
        ws.cell(row=row_num, column=1, value=f"=ROW()-9").alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.cell(row=row_num, column=2, value=item[0]).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        ws.cell(row=row_num, column=3, value=item[1]).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        ws.cell(row=row_num, column=4, value=item[2]).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        ws.cell(row=row_num, column=5, value=item[3]).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        ws.cell(row=row_num, column=6, value=item[4]).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        ws.cell(row=row_num, column=7, value=item[5]).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        # Công thức Excel động đánh giá ĐẠT / KHÔNG ĐẠT từ cột G
        ws.cell(
            row=row_num,
            column=8,
            value=f'=IF(G{row_num}=1, "ĐẠT", "KHÔNG ĐẠT")',
        ).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=9, value=item[6]).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )

        for col_idx in range(1, 10):
            c = ws.cell(row=row_num, column=col_idx)
            c.font = bold_font if col_idx in (1, 7, 8) else regular_font
            c.border = grid_border
            c.fill = white_fill
        ws.row_dimensions[row_num].height = 28

    # Thiết lập độ rộng cột chuẩn
    col_widths = [6, 20, 38, 24, 30, 30, 16, 16, 22]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    target_xlsx = (
        Path(__file__).resolve().parent.parent
        / "checklists"
        / "Bang_Kiem_Tra_Khao_Sat_Giam_Sat_Hien_Truong.xlsx"
    )
    saved = create_dynamic_field_checklist(target_xlsx)
    print(f"Đã tạo Bảng kiểm tra hiện trường Excel động tại: {saved}")
